"""Build the Excel business review workbook from the warehouse.

The workbook is generated rather than hand-built so it can be rebuilt
identically whenever the data changes, and so every formula is reviewable as
code. Every number a reader sees in the workbook is a live Excel formula over
the raw sheets: change a raw value and the whole workbook recalculates.

Scope: customers who signed up January-June 2025 and all of their orders. That
is ~3,400 customers and ~8,700 orders, which is comfortable for Excel. The full
warehouse (31,000 orders, 60,000 line items) is where Excel starts to strain,
which is exactly why projects 1 and 4 use SQL. The workbook says so on its
README sheet.

Usage (from the repository root, after `make warehouse`):
    python 00-excel-business-review/build_workbook.py --db warehouse.duckdb
"""

from __future__ import annotations

import argparse
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "business_review.xlsx"

SIGNUP_FROM, SIGNUP_TO = "2025-01-01", "2025-06-30"

# ---- styles ----------------------------------------------------------------
FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color="1F3A5F")
F_SUB = Font(name=FONT, size=11, bold=True, color="1F3A5F")
F_INPUT = Font(name=FONT, size=10, color="0000FF")          # hardcoded inputs
F_LINK = Font(name=FONT, size=10, color="008000")           # cross-sheet links
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_KPI = Font(name=FONT, size=20, bold=True, color="1F3A5F")
FILL_HEAD = PatternFill("solid", fgColor="1F3A5F")
FILL_RAW = PatternFill("solid", fgColor="7F7F7F")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_KPI = PatternFill("solid", fgColor="EAF1FB")
FILL_BAND = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(wrap_text=True, vertical="top")

NGN = '"₦"#,##0;("₦"#,##0);"-"'
NGN2 = '"₦"#,##0.00;("₦"#,##0.00);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "yyyy-mm-dd"
MONTH = "mmm yyyy"


def header(ws, row, labels, fill=FILL_HEAD, widths=None):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.alignment, c.border = F_HEAD, fill, CENTER, BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def body_font(ws, min_row, max_row, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col):
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = F_BODY


# ---- data ------------------------------------------------------------------

def load(con):
    con.execute((ROOT / "01-sql-retail-warehouse" / "sql" / "01_staging.sql").read_text())
    customers = con.execute(f"""
        SELECT customer_id, email, signup_date, coalesce(country, '') AS country,
               acquisition_channel
        FROM raw_customers
        WHERE CAST(signup_date AS DATE) BETWEEN DATE '{SIGNUP_FROM}' AND DATE '{SIGNUP_TO}'
        ORDER BY CAST(customer_id AS BIGINT)
    """).fetchall()
    # Orders arrive with revenue already rolled up from line items, de-duplicated
    # and with the x100 price typos excluded -- that item-level cleaning is the
    # SQL project's job (stg_order_items). Everything from here on is Excel's.
    orders = con.execute(f"""
        WITH items AS (
            SELECT order_id, sum(gross_revenue) AS gross, sum(return_value) AS ret
            FROM stg_order_items GROUP BY order_id
        )
        SELECT o.order_id, o.customer_id, o.order_ts, o.status, o.channel,
               coalesce(o.discount_code, '') AS discount_code,
               round(coalesce(i.gross, 0), 2) AS gross_revenue,
               round(coalesce(i.ret, 0), 2)   AS return_value
        FROM raw_orders o
        JOIN raw_customers c USING (customer_id)
        LEFT JOIN items i ON i.order_id = CAST(o.order_id AS BIGINT)
        WHERE CAST(c.signup_date AS DATE) BETWEEN DATE '{SIGNUP_FROM}' AND DATE '{SIGNUP_TO}'
        ORDER BY CAST(o.order_id AS BIGINT)
    """).fetchall()
    return customers, orders


# ---- sheets ----------------------------------------------------------------

def sheet_readme(wb, n_cust, n_orders):
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 92
    ws["B2"] = "Business Review Workbook"
    ws["B2"].font = F_TITLE
    ws["B3"] = ("Monthly revenue, retention and customer segments for the January-June 2025 "
                "signup cohorts, built entirely from Excel formulas.")
    ws["B3"].font = F_NOTE
    rows = [
        ("QUESTION", "Of the customers we acquired in H1 2025, how many are still buying, "
                     "who is worth saving, and is growth coming from new or returning customers?"),
        ("DATA", f"{n_cust:,} customers and {n_orders:,} orders, exported from the portfolio "
                 "warehouse (synthetic, generated deterministically - see the repository README). "
                 "Raw sheets are pasted exactly as exported, warts included."),
        ("HOW IT FLOWS", "Raw_Customers + Raw_Orders  ->  Customers + Clean_Orders (cleaning "
                         "formulas)  ->  Monthly, Cohorts, Segments, Channels (analysis)  ->  "
                         "Dashboard (one page)  ->  Checks (reconciliation against SQL)"),
        ("COLOUR KEY", "Blue text = a hardcoded input you may change. Yellow fill = the one "
                       "control on the Dashboard (channel picker). Green text = a value pulled "
                       "from another sheet. Black = a formula. Grey header = raw data, never "
                       "edit by hand."),
        ("CLEANING DONE HERE", "1) status has 8 spellings -> LOWER(TRIM()) gives 5 states. "
                               "2) order_ts arrives in two formats (ISO and DD/MM/YYYY) -> one "
                               "formula parses both. 3) emails have stray case and spaces -> "
                               "LOWER(TRIM()). 4) blank countries -> 'unknown' so they survive "
                               "every summary. 5) internal test accounts (@test.internal) are "
                               "flagged and excluded from analysis, not deleted."),
        ("CLEANING DONE UPSTREAM", "Order revenue arrives already rolled up from line items, "
                                   "with duplicate rows removed and x100 price typos excluded. "
                                   "That item-level cleaning lives in the SQL project "
                                   "(01-sql-retail-warehouse/sql/01_staging.sql)."),
        ("REVENUE RULE", "An order counts as revenue only if its cleaned status is 'completed' "
                         "or 'shipped', and the customer is not a test account. Net revenue = "
                         "gross revenue - returns. Defined once, in Clean_Orders column K."),
        ("TRUST", "The Checks sheet compares 12 headline numbers in this workbook against the "
                  "same numbers computed independently in SQL. All must show MATCH."),
        ("WHY NOT EXCEL FOR EVERYTHING", "This slice is ~8,700 orders. The full warehouse is "
                                         "31,000 orders and 60,000 line items; distinct-count "
                                         "formulas like the ones in Clean_Orders column N slow "
                                         "down quadratically with row count. That is the point "
                                         "where the work moves to SQL - see projects 1 and 4."),
        ("PIVOT TABLES", "openpyxl cannot write native PivotTables, so this workbook uses "
                         "formulas. The Pivot_Guide sheet gives click-by-click steps to rebuild "
                         "the Monthly and Segments sheets as PivotTables, and to load the raw "
                         "data with Power Query."),
    ]
    r = 5
    for label, text in rows:
        ws.cell(row=r, column=2, value=label).font = F_SUB
        c = ws.cell(row=r, column=3, value=text)
        c.font, c.alignment = F_BODY, WRAP
        ws.row_dimensions[r].height = max(30, 15 * (len(text) // 88 + 1))
        r += 1
    ws.sheet_view.showGridLines = False


def sheet_raw_customers(wb, customers):
    ws = wb.create_sheet("Raw_Customers")
    header(ws, 1, ["customer_id", "email", "signup_date", "country", "acquisition_channel"],
           fill=FILL_RAW, widths=[13, 34, 13, 10, 20])
    for i, (cid, email, signup, country, chan) in enumerate(customers, 2):
        ws.cell(row=i, column=1, value=int(cid))
        ws.cell(row=i, column=2, value=email)          # left messy on purpose
        ws.cell(row=i, column=3, value=signup)         # ISO text, as exported
        ws.cell(row=i, column=4, value=country)
        ws.cell(row=i, column=5, value=chan)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{len(customers) + 1}"
    return len(customers) + 1


def sheet_raw_orders(wb, orders):
    ws = wb.create_sheet("Raw_Orders")
    header(ws, 1, ["order_id", "customer_id", "order_ts", "status", "channel",
                   "discount_code", "gross_revenue", "return_value"],
           fill=FILL_RAW, widths=[11, 13, 19, 12, 13, 14, 15, 13])
    for i, row in enumerate(orders, 2):
        oid, cid, ts, status, chan, disc, gross, ret = row
        ws.cell(row=i, column=1, value=int(oid))
        ws.cell(row=i, column=2, value=int(cid))
        ws.cell(row=i, column=3, value=ts)             # text, two formats
        ws.cell(row=i, column=4, value=status)         # eight spellings
        ws.cell(row=i, column=5, value=chan)
        ws.cell(row=i, column=6, value=disc)
        ws.cell(row=i, column=7, value=float(gross)).number_format = NGN2
        ws.cell(row=i, column=8, value=float(ret)).number_format = NGN2
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(orders) + 1}"
    return len(orders) + 1


def sheet_customers(wb, last_cust, last_ord):
    """One row per customer: cleaned fields plus lifetime metrics and RFM."""
    ws = wb.create_sheet("Customers")
    cols = ["customer_id", "email_clean", "signup_date", "signup_month", "country",
            "acquisition_channel", "is_test", "orders", "net_revenue", "last_order_date",
            "recency_days", "R", "F", "M", "segment"]
    header(ws, 1, cols, widths=[13, 32, 12, 12, 10, 18, 9, 9, 14, 14, 11, 6, 6, 6, 20])
    n = last_cust
    co = f"Clean_Orders!$B$2:$B${last_ord}"         # customer_id
    inana = f"Clean_Orders!$L$2:$L${last_ord}"      # in_analysis flag
    net = f"Clean_Orders!$K$2:$K${last_ord}"        # net revenue
    odate = f"Clean_Orders!$C$2:$C${last_ord}"      # order date
    rec = f"$K$2:$K${n}"
    frq = f"$H$2:$H${n}"
    mon = f"$I$2:$I${n}"
    for r in range(2, n + 1):
        ws[f"A{r}"] = f"=Raw_Customers!A{r}"
        ws[f"B{r}"] = f"=LOWER(TRIM(Raw_Customers!B{r}))"
        # ISO text -> a real date. DATE(LEFT, MID, RIGHT) is locale-proof;
        # DATEVALUE reads the same text differently on a UK vs a US machine.
        ws[f"C{r}"] = (f"=DATE(VALUE(LEFT(Raw_Customers!C{r},4)),"
                       f"VALUE(MID(Raw_Customers!C{r},6,2)),VALUE(RIGHT(Raw_Customers!C{r},2)))")
        ws[f"D{r}"] = f"=DATE(YEAR(C{r}),MONTH(C{r}),1)"
        ws[f"E{r}"] = f'=IF(TRIM(Raw_Customers!D{r})="","unknown",TRIM(Raw_Customers!D{r}))'
        ws[f"F{r}"] = f"=TRIM(Raw_Customers!E{r})"
        ws[f"G{r}"] = f'=ISNUMBER(SEARCH("@test.internal",B{r}))'
        ws[f"H{r}"] = f"=COUNTIFS({co},A{r},{inana},TRUE)"
        ws[f"I{r}"] = f"=SUMIFS({net},{co},A{r},{inana},TRUE)"
        ws[f"J{r}"] = f'=IF(H{r}=0,"",_xlfn.MAXIFS({odate},{co},A{r},{inana},TRUE))'
        ws[f"K{r}"] = f'=IF(H{r}=0,"",Dashboard!$C$4-J{r})'
        # Quintile scores without NTILE. Score = 1 + how many buyers are worse,
        # as a share of all buyers, in fifths. Ties always get the same score,
        # so this is deterministic -- unlike NTILE, which splits ties arbitrarily
        # (the bug documented in the SQL project). COUNT() counts numbers only,
        # so non-buyers ("" cells) drop out of the denominator automatically.
        ws[f"L{r}"] = f'=IF(H{r}=0,"",MIN(5,1+INT(5*COUNTIFS({rec},">"&K{r})/COUNT({rec}))))'
        ws[f"M{r}"] = (f'=IF(H{r}=0,"",MIN(5,1+INT(5*COUNTIFS({frq},"<"&H{r},'
                       f'{frq},">0")/COUNTIFS({frq},">0"))))')
        ws[f"N{r}"] = (f'=IF(H{r}=0,"",MIN(5,1+INT(5*COUNTIFS({mon},"<"&I{r},'
                       f'{frq},">0")/COUNTIFS({frq},">0"))))')
        ws[f"O{r}"] = (f'=IF(OR(G{r},H{r}=0),"",'
                       f'IF(AND(L{r}>=4,M{r}>=4),"champion",'
                       f'IF(AND(L{r}>=4,M{r}<=2),"promising new",'
                       f'IF(AND(L{r}<=2,M{r}>=4),"at risk - was loyal",'
                       f'IF(AND(L{r}<=2,M{r}<=2),"lapsed",'
                       f'IF(N{r}=5,"big spender","steady"))))))')
        for col, fmt in (("C", DATE), ("D", MONTH), ("I", NGN), ("J", DATE), ("K", INT)):
            ws[f"{col}{r}"].number_format = fmt
    for r in range(2, n + 1):
        ws[f"A{r}"].font = F_LINK
    body_font(ws, 2, n, len(cols))
    for r in range(2, n + 1):
        ws[f"A{r}"].font = F_LINK
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:O{n}"
    ws["Q1"] = "How RFM is scored here"
    ws["Q1"].font = F_SUB
    ws["Q2"] = ("R, F, M are 1-5. A score is 1 + the share of buyers who are worse on that "
                "measure, in fifths. Ties share a score, so results can differ slightly from "
                "the SQL project, which uses NTILE plus a tiebreaker.")
    ws["Q2"].font, ws["Q2"].alignment = F_NOTE, WRAP
    ws.column_dimensions["Q"].width = 50
    ws.row_dimensions[2].height = 60


def sheet_clean_orders(wb, last_ord, last_cust):
    """One row per order, cleaned, with every flag the analysis sheets need."""
    ws = wb.create_sheet("Clean_Orders")
    cols = ["order_id", "customer_id", "order_date", "order_month", "status_clean",
            "recognised", "is_test", "signup_month", "months_since_signup",
            "acquisition_channel", "net_revenue", "in_analysis", "first_order_month",
            "first_in_month", "is_new_customer_order"]
    header(ws, 1, cols, widths=[11, 13, 12, 12, 12, 11, 9, 12, 11, 18, 14, 11, 13, 11, 13])
    ccid = f"Customers!$A$2:$A${last_cust}"
    n = last_ord
    for r in range(2, n + 1):
        ws[f"A{r}"] = f"=Raw_Orders!A{r}"
        ws[f"B{r}"] = f"=Raw_Orders!B{r}"
        # Two timestamp formats arrive. Character 5 is "-" only in the ISO form
        # (2025-03-04 ...); the other is DD/MM/YYYY. A plain DATEVALUE would
        # silently misread or reject one of them -- the same 2% of rows a naive
        # SQL CAST drops.
        ws[f"C{r}"] = (f'=IF(MID(Raw_Orders!C{r},5,1)="-",'
                       f"DATE(VALUE(LEFT(Raw_Orders!C{r},4)),VALUE(MID(Raw_Orders!C{r},6,2)),"
                       f"VALUE(MID(Raw_Orders!C{r},9,2))),"
                       f"DATE(VALUE(MID(Raw_Orders!C{r},7,4)),VALUE(MID(Raw_Orders!C{r},4,2)),"
                       f"VALUE(LEFT(Raw_Orders!C{r},2))))")
        ws[f"D{r}"] = f"=DATE(YEAR(C{r}),MONTH(C{r}),1)"
        ws[f"E{r}"] = f"=LOWER(TRIM(Raw_Orders!D{r}))"
        ws[f"F{r}"] = f'=OR(E{r}="completed",E{r}="shipped")'
        ws[f"G{r}"] = f"=INDEX(Customers!$G$2:$G${last_cust},MATCH(B{r},{ccid},0))"
        ws[f"H{r}"] = f"=INDEX(Customers!$D$2:$D${last_cust},MATCH(B{r},{ccid},0))"
        ws[f"I{r}"] = f"=(YEAR(D{r})-YEAR(H{r}))*12+MONTH(D{r})-MONTH(H{r})"
        ws[f"J{r}"] = f"=INDEX(Customers!$F$2:$F${last_cust},MATCH(B{r},{ccid},0))"
        ws[f"K{r}"] = f"=IF(F{r},Raw_Orders!G{r}-Raw_Orders!H{r},0)"
        ws[f"L{r}"] = f"=AND(F{r},NOT(G{r}))"
        ws[f"M{r}"] = (f'=_xlfn.MINIFS($D$2:$D${n},$B$2:$B${n},B{r},$L$2:$L${n},TRUE)')
        # Distinct counting with COUNTIFS. The range grows one row at a time
        # ($B$2:B2, $B$2:B3, ...), so the formula asks "is this the first
        # analysis order I have seen for this customer in this month?". Summing
        # the TRUEs gives distinct buyers per month. This is the expanding-range
        # trick, and it is also why Excel slows down on big tables: row 8,000
        # scans 8,000 rows.
        ws[f"N{r}"] = (f"=AND(L{r},COUNTIFS($B$2:B{r},B{r},$D$2:D{r},D{r},"
                       f"$L$2:L{r},TRUE)=1)")
        ws[f"O{r}"] = f"=AND(L{r},D{r}=M{r})"
        for col, fmt in (("C", DATE), ("D", MONTH), ("H", MONTH), ("K", NGN), ("M", MONTH)):
            ws[f"{col}{r}"].number_format = fmt
    body_font(ws, 2, n, len(cols))
    for r in range(2, n + 1):
        ws[f"A{r}"].font = F_LINK
        ws[f"B{r}"].font = F_LINK
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:O{n}"


def sheet_monthly(wb, last_ord):
    ws = wb.create_sheet("Monthly")
    ws["A1"] = "Monthly revenue - H1 2025 signup cohorts"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Every figure is SUMIFS / COUNTIFS over Clean_Orders where in_analysis = TRUE. "
                "Column J follows the channel picked on the Dashboard.")
    ws["A2"].font = F_NOTE
    cols = ["month", "net_revenue", "orders", "buyers", "new_customer_revenue",
            "returning_revenue", "pct_returning", "avg_order_value", "mom_growth",
            "revenue_selected_channel"]
    header(ws, 4, cols, widths=[12, 15, 9, 9, 16, 16, 12, 14, 12, 18])
    o = lambda col: f"Clean_Orders!${col}$2:${col}${last_ord}"  # noqa: E731
    for i in range(12):
        r = 5 + i
        ws[f"A{r}"] = f"=DATE(2025,{i + 1},1)"
        ws[f"A{r}"].number_format = MONTH
        ws[f"B{r}"] = f"=SUMIFS({o('K')},{o('D')},A{r},{o('L')},TRUE)"
        ws[f"C{r}"] = f"=COUNTIFS({o('D')},A{r},{o('L')},TRUE)"
        ws[f"D{r}"] = f"=COUNTIFS({o('D')},A{r},{o('N')},TRUE)"
        ws[f"E{r}"] = f"=SUMIFS({o('K')},{o('D')},A{r},{o('O')},TRUE)"
        ws[f"F{r}"] = f"=B{r}-E{r}"
        ws[f"G{r}"] = f"=IFERROR(F{r}/B{r},0)"
        ws[f"H{r}"] = f"=IFERROR(B{r}/C{r},0)"
        ws[f"I{r}"] = "" if i == 0 else f'=IFERROR(B{r}/B{r - 1}-1,"")'
        ws[f"J{r}"] = (f'=IF(Dashboard!$C$6="All",B{r},'
                       f"SUMIFS({o('K')},{o('D')},A{r},{o('L')},TRUE,{o('J')},Dashboard!$C$6))")
        for col, fmt in (("B", NGN), ("C", INT), ("D", INT), ("E", NGN), ("F", NGN),
                         ("G", PCT), ("H", NGN), ("I", PCT), ("J", NGN)):
            ws[f"{col}{r}"].number_format = fmt
    r = 17
    ws[f"A{r}"] = "Total"
    ws[f"A{r}"].font = F_BOLD
    for col in "BCEF":
        ws[f"{col}{r}"] = f"=SUM({col}5:{col}16)"
        ws[f"{col}{r}"].font = F_BOLD
        ws[f"{col}{r}"].number_format = NGN if col in "BEF" else INT
    ws[f"G{r}"] = f"=IFERROR(F{r}/B{r},0)"
    ws[f"G{r}"].number_format = PCT
    ws[f"H{r}"] = f"=IFERROR(B{r}/C{r},0)"
    ws[f"H{r}"].number_format = NGN
    ws["A19"] = ("Buyers are not totalled: the same customer buys in many months, so adding "
                 "monthly buyers double-counts. Distinct buyers for the year are on Checks.")
    ws["A19"].font = F_NOTE
    body_font(ws, 5, 17, 10)
    ws.freeze_panes = "B5"


def sheet_cohorts(wb, last_ord, last_cust):
    ws = wb.create_sheet("Cohorts")
    ws["A1"] = "Cohort retention - share of each signup cohort still buying"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Rows = signup month. Columns = months since signup. A cell is blank when "
                "the cohort has not had that many months yet (censored), never zero.")
    ws["A2"].font = F_NOTE
    labels = ["signup_month", "cohort_size"] + [f"M{k}" for k in range(12)]
    header(ws, 4, labels, widths=[13, 11] + [7] * 12)
    o = lambda col: f"Clean_Orders!${col}$2:${col}${last_ord}"  # noqa: E731
    c = lambda col: f"Customers!${col}$2:${col}${last_cust}"  # noqa: E731
    ws["P4"] = "data ends"
    ws["P4"].font = F_BOLD
    ws["P5"] = f"=MAX({o('D')})"
    ws["P5"].number_format = MONTH
    for i in range(6):
        r = 5 + i
        ws[f"A{r}"] = f"=DATE(2025,{i + 1},1)"
        ws[f"A{r}"].number_format = MONTH
        ws[f"B{r}"] = f"=COUNTIFS({c('D')},A{r},{c('G')},FALSE)"
        ws[f"B{r}"].number_format = INT
        for k in range(12):
            col = get_column_letter(3 + k)
            months_available = f"((YEAR($P$5)-YEAR($A{r}))*12+MONTH($P$5)-MONTH($A{r}))"
            ws[f"{col}{r}"] = (f'=IF({k}>{months_available},"",'
                               f"COUNTIFS({o('H')},$A{r},{o('I')},{k},{o('N')},TRUE)/$B{r})")
            ws[f"{col}{r}"].number_format = '0.0%'
    body_font(ws, 5, 10, 14)
    ws.conditional_formatting.add(
        "C5:N10",
        ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                       end_type="max", end_color="2A78D6"))
    ws["A12"] = "Average over cohorts that have the month"
    ws["A12"].font = F_BOLD
    for k in range(12):
        col = get_column_letter(3 + k)
        ws[f"{col}12"] = f'=IFERROR(AVERAGE({col}5:{col}10),"")'
        ws[f"{col}12"].number_format = '0.0%'
        ws[f"{col}12"].font = F_BOLD
    ws["A13"] = ("Reading tip: month 0 is a partial month (signups late in a month have "
                 "little time to buy), so month 1 can be higher than month 0.")
    ws["A13"].font = F_NOTE


def sheet_segments(wb, last_cust):
    ws = wb.create_sheet("Segments")
    ws["A1"] = "RFM segments - who to protect, who to win back"
    ws["A1"].font = F_TITLE
    header(ws, 3, ["segment", "customers", "avg_orders", "avg_revenue",
                   "avg_recency_days", "segment_revenue", "pct_of_revenue", "action"],
           widths=[22, 11, 11, 14, 15, 16, 14, 16])
    c = lambda col: f"Customers!${col}$2:${col}${last_cust}"  # noqa: E731
    segs = [("champion", "protect"), ("steady", "monitor"),
            ("at risk - was loyal", "act now"), ("big spender", "monitor"),
            ("lapsed", "low priority"), ("promising new", "nurture")]
    for i, (seg, action) in enumerate(segs):
        r = 4 + i
        ws[f"A{r}"] = seg
        ws[f"B{r}"] = f"=COUNTIFS({c('O')},A{r})"
        ws[f"C{r}"] = f"=IFERROR(AVERAGEIFS({c('H')},{c('O')},A{r}),0)"
        ws[f"D{r}"] = f"=IFERROR(AVERAGEIFS({c('I')},{c('O')},A{r}),0)"
        ws[f"E{r}"] = f"=IFERROR(AVERAGEIFS({c('K')},{c('O')},A{r}),0)"
        ws[f"F{r}"] = f"=SUMIFS({c('I')},{c('O')},A{r})"
        ws[f"G{r}"] = f"=IFERROR(F{r}/$F$10,0)"
        ws[f"H{r}"] = action
        for col, fmt in (("B", INT), ("C", "0.00"), ("D", NGN), ("E", "#,##0"),
                         ("F", NGN), ("G", PCT)):
            ws[f"{col}{r}"].number_format = fmt
    ws["A10"] = "Total"
    ws["A10"].font = F_BOLD
    ws["B10"] = "=SUM(B4:B9)"
    ws["F10"] = "=SUM(F4:F9)"
    ws["G10"] = "=SUM(G4:G9)"
    for col, fmt in (("B", INT), ("F", NGN), ("G", PCT)):
        ws[f"{col}10"].number_format = fmt
        ws[f"{col}10"].font = F_BOLD
    body_font(ws, 4, 9, 8)
    ws.conditional_formatting.add(
        "A4:H9", FormulaRule(formula=['$H4="act now"'],
                             fill=PatternFill("solid", fgColor="FDE9D9")))


def sheet_channels(wb, last_cust):
    ws = wb.create_sheet("Channels")
    ws["A1"] = "Acquisition channel quality"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Revenue per acquired customer, by the channel that brought them in. "
                "No spend data in this slice, so this ranks quality, not return on spend.")
    ws["A2"].font = F_NOTE
    header(ws, 4, ["channel", "customers", "buyers", "conversion", "net_revenue",
                   "revenue_per_customer", "revenue_per_buyer"],
           widths=[16, 11, 10, 12, 15, 19, 17])
    c = lambda col: f"Customers!${col}$2:${col}${last_cust}"  # noqa: E731
    for i, ch in enumerate(["email", "referral", "organic", "direct",
                            "paid_search", "paid_social"]):
        r = 5 + i
        ws[f"A{r}"] = ch
        ws[f"B{r}"] = f"=COUNTIFS({c('F')},A{r},{c('G')},FALSE)"
        ws[f"C{r}"] = f'=COUNTIFS({c("F")},A{r},{c("G")},FALSE,{c("H")},">0")'
        ws[f"D{r}"] = f"=IFERROR(C{r}/B{r},0)"
        ws[f"E{r}"] = f"=SUMIFS({c('I')},{c('F')},A{r},{c('G')},FALSE)"
        ws[f"F{r}"] = f"=IFERROR(E{r}/B{r},0)"
        ws[f"G{r}"] = f"=IFERROR(E{r}/C{r},0)"
        for col, fmt in (("B", INT), ("C", INT), ("D", PCT), ("E", NGN),
                         ("F", NGN), ("G", NGN)):
            ws[f"{col}{r}"].number_format = fmt
    body_font(ws, 5, 10, 7)
    ws.conditional_formatting.add(
        "F5:F10", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                 end_type="max", end_color="1BAF7A"))


def sheet_dashboard(wb, last_ord, last_cust):
    ws = wb.create_sheet("Dashboard", 1)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJ", [2, 26, 18, 3, 26, 18, 3, 26, 18, 3], strict=True):
        ws.column_dimensions[col].width = w
    ws["B2"] = "H1 2025 cohorts - business review"
    ws["B2"].font = F_TITLE
    ws["B4"] = "As-of date (last order)"
    o = lambda col: f"Clean_Orders!${col}$2:${col}${last_ord}"  # noqa: E731
    ws["C4"] = f"=_xlfn.MAXIFS({o('C')},{o('L')},TRUE)"
    ws["C4"].number_format = DATE
    ws["B6"] = "Channel filter (pick one)"
    ws["C6"] = "All"
    ws["C6"].font, ws["C6"].fill = F_INPUT, FILL_INPUT
    ws["C6"].comment = Comment("Change this to filter the monthly chart and the "
                               "'selected channel' KPI. 'All' shows every channel.", "workbook")
    dv = DataValidation(type="list", allow_blank=False,
                        formula1='"All,email,referral,organic,direct,paid_search,paid_social"')
    ws.add_data_validation(dv)
    dv.add("C6")

    kpis = [
        ("B9", "Net revenue, Jan-Dec 2025", "=Monthly!B17", NGN),
        ("E9", "Revenue from returning customers", "=Monthly!G17", PCT),
        ("H9", "Avg order value", "=Monthly!H17", NGN),
        ("B13", "Customers acquired (H1)", "=SUM(Cohorts!B5:B10)", INT),
        ("E13", "Average month-1 retention", "=Cohorts!D12", PCT),
        ("H13", "At-risk revenue (share)", "=Segments!G6", PCT),
        ("B17", "Selected channel revenue", "=SUM(Monthly!J5:J16)", NGN),
        ("E17", "Best channel, revenue / customer",
         "=INDEX(Channels!A5:A10,MATCH(MAX(Channels!F5:F10),Channels!F5:F10,0))", None),
        ("H17", "Buyers who ordered again", "=Checks!C18", PCT),
    ]
    for cell, label, formula, fmt in kpis:
        col, row = cell[0], int(cell[1:])
        nxt = chr(ord(col) + 1)
        ws[f"{col}{row}"] = label
        ws[f"{col}{row}"].font = F_NOTE
        ws.merge_cells(f"{col}{row + 1}:{nxt}{row + 2}")
        v = ws[f"{col}{row + 1}"]
        v.value = formula
        v.font, v.fill, v.alignment = F_KPI, FILL_KPI, CENTER
        if fmt:
            v.number_format = fmt

    line = LineChart()
    line.title = "Net revenue by month (selected channel)"
    line.y_axis.title = "₦"
    line.height, line.width = 7.5, 13.5
    line.add_data(Reference(wb["Monthly"], min_col=10, min_row=4, max_row=16),
                  titles_from_data=True)
    line.set_categories(Reference(wb["Monthly"], min_col=1, min_row=5, max_row=16))
    line.legend = None
    # Smoothing draws curves between points that the data never contained, so
    # a dip can look like it bottomed out between months. Straight segments only.
    for series in line.series:
        series.smooth = False
        series.graphicalProperties.line.solidFill = "2A78D6"
        series.graphicalProperties.line.width = 28575  # 2.25pt in EMU
    ws.add_chart(line, "B21")

    bar = BarChart()
    bar.type, bar.grouping, bar.overlap = "col", "stacked", 100
    bar.title = "New vs returning revenue"
    bar.height, bar.width = 7.5, 13.5
    bar.add_data(Reference(wb["Monthly"], min_col=5, max_col=6, min_row=4, max_row=16),
                 titles_from_data=True)
    bar.set_categories(Reference(wb["Monthly"], min_col=1, min_row=5, max_row=16))
    # Same colours as the web dashboard: colour follows the entity everywhere.
    bar.series[0].graphicalProperties.solidFill = "2A78D6"   # new
    bar.series[1].graphicalProperties.solidFill = "EB6834"   # returning
    bar.legend.position = "b"
    ws.add_chart(bar, "F21")

    seg = BarChart()
    seg.type = "bar"
    seg.title = "Revenue share by segment"
    seg.height, seg.width = 7.5, 13.5
    seg.add_data(Reference(wb["Segments"], min_col=7, min_row=3, max_row=9),
                 titles_from_data=True)
    seg.set_categories(Reference(wb["Segments"], min_col=1, min_row=4, max_row=9))
    seg.legend = None
    seg.series[0].graphicalProperties.solidFill = "2A78D6"
    ws.add_chart(seg, "B37")


def sheet_checks(wb, expected, last_ord, last_cust):
    """Reconcile workbook numbers against the same numbers computed in SQL."""
    ws = wb.create_sheet("Checks")
    ws["A1"] = "Checks - does this workbook agree with SQL?"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Blue = expected value computed independently in DuckDB by build_workbook.py "
                "(Source: warehouse.duckdb, same slice and rules). Black = this workbook's "
                "formula. Tolerance 0.5 for currency rounding.")
    ws["A2"].font = F_NOTE
    header(ws, 4, ["check", "expected (SQL)", "workbook", "status"], widths=[44, 18, 18, 10])
    o = lambda col: f"Clean_Orders!${col}$2:${col}${last_ord}"  # noqa: E731
    c = lambda col: f"Customers!${col}$2:${col}${last_cust}"  # noqa: E731
    checks = [
        ("Customers in slice", expected["customers"], f"=COUNTA({c('A')})", INT),
        ("Orders in slice", expected["orders"], f"=COUNTA({o('A')})", INT),
        ("Test accounts", expected["test_accounts"], f"=COUNTIFS({c('G')},TRUE)", INT),
        ("Orders counted as revenue", expected["analysis_orders"],
         f"=COUNTIFS({o('L')},TRUE)", INT),
        ("Net revenue", expected["net_revenue"], f"=SUMIFS({o('K')},{o('L')},TRUE)", NGN),
        ("Timestamps parsed (all must be dates)", expected["orders"], f"=COUNT({o('C')})", INT),
        ("Distinct status values after cleaning", expected["statuses"],
         f"=SUMPRODUCT(1/COUNTIF({o('E')},{o('E')}))", INT),
        ("Distinct buyers", expected["buyers"], f"=COUNTIFS({c('H')},\">0\")", INT),
        ("Buyers with 2+ orders", expected["repeat_buyers"], f"=COUNTIFS({c('H')},\">1\")", INT),
        ("Revenue from first-month orders", expected["new_revenue"],
         f"=SUMIFS({o('K')},{o('O')},TRUE)", NGN),
        ("Month-1 retention, Jan 2025 cohort", expected["jan_m1"], "=Cohorts!D5", PCT),
        ("Blank (unknown) countries", expected["unknown_country"],
         f'=COUNTIFS({c("E")},"unknown")', INT),
    ]
    for i, (label, exp, formula, fmt) in enumerate(checks):
        r = 5 + i
        ws[f"A{r}"] = label
        ws[f"B{r}"] = exp
        ws[f"B{r}"].font = F_INPUT
        ws[f"C{r}"] = formula
        ws[f"D{r}"] = (f'=IF(ABS(B{r}-C{r})<=IF(ABS(B{r})<1,0.0005,0.5),"MATCH","CHECK")')
        for col in "BC":
            ws[f"{col}{r}"].number_format = fmt
    # Derived figures sit BELOW the check table. An earlier version wrote this
    # row on top of the last check, which silently replaced it -- the Checks
    # sheet caught its own bug, which is the argument for having one.
    last = 4 + len(checks)
    ws[f"A{last + 2}"] = "Share of buyers who ordered more than once"
    ws[f"C{last + 2}"] = "=IFERROR(C13/C12,0)"
    ws[f"C{last + 2}"].number_format = PCT
    body_font(ws, 5, last + 2, 4)
    for r in range(5, last + 1):
        ws[f"B{r}"].font = F_INPUT
    rng = f"D5:D{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"MATCH"'],
                                  fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CHECK"'],
                                  fill=PatternFill("solid", fgColor="FFC7CE")))
    ws[f"A{last + 1}"] = "All checks"
    ws[f"A{last + 1}"].font = F_BOLD
    ws[f"D{last + 1}"] = f'=IF(COUNTIF({rng},"MATCH")={len(checks)},"ALL MATCH","FIX FIRST")'
    ws[f"D{last + 1}"].font = F_BOLD
    ws[f"A{last + 4}"] = ("Why this sheet exists: two independent calculations agreeing "
                          "is the cheapest "
                 "real proof either one is right. If any row says CHECK, fix it before "
                 "showing anyone a number.")
    ws[f"A{last + 4}"].font = F_NOTE


def sheet_pivot_guide(wb):
    ws = wb.create_sheet("Pivot_Guide")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    ws["B1"] = "Rebuild it with PivotTables and Power Query"
    ws["B1"].font = F_TITLE
    steps = [
        ("A. Monthly revenue as a PivotTable", None),
        ("1", "Click any cell in Clean_Orders. Insert > PivotTable > New Worksheet > OK."),
        ("2", "Drag in_analysis to Filters and set it to TRUE."),
        ("3", "Drag order_month to Rows. (If Excel groups it into Years/Quarters, right-click > "
              "Ungroup, or keep Months only.)"),
        ("4", "Drag net_revenue to Values. It shows 'Sum of net_revenue'. "
              "Number format: currency."),
        ("5", "Drag order_id to Values too, then Value Field Settings > Count. That is orders."),
        ("6", "Compare with the Monthly sheet. The numbers must match. If they do not, a filter "
              "is wrong - check in_analysis first."),
        ("B. Distinct buyers (the PivotTable trap)", None),
        ("7", "Count of customer_id counts ORDERS, not customers. To count distinct customers, "
              "tick 'Add this data to the Data Model' when creating the PivotTable, then use "
              "Value Field Settings > Distinct Count."),
        ("C. Segments as a PivotTable", None),
        ("8", "From the Customers sheet: segment to Rows, customer_id to Values (Count), "
              "net_revenue to Values (Sum), then Show Values As > % of Grand Total."),
        ("D. Load raw data with Power Query (the professional way)", None),
        ("9", "Data > Get Data > From File > From Text/CSV, pick orders.csv from data/raw."),
        ("10", "Transform Data. Select status > Transform > Format > Trim, then lowercase."),
        ("11", "order_ts: Split Column by delimiter ' ', keep the date part. Then use "
               "'Using Locale' to convert the DD/MM/YYYY rows. Power Query records each step."),
        ("12", "Close & Load. Next month, drop in the new CSV and click Refresh: every step "
               "re-runs. That repeatability is the reason Power Query beats copy-paste."),
        ("E. When to stop using Excel", None),
        ("13", "If a sheet takes more than a few seconds to recalculate, or you are past "
               "~100k rows, or several people need the same numbers, move the logic to SQL. "
               "Excel is the presentation and exploration layer, not the warehouse."),
    ]
    r = 3
    for num, text in steps:
        if text is None:
            ws[f"B{r}"] = num
            ws[f"B{r}"].font = F_SUB
            r += 1
            continue
        ws[f"A{r}"] = num
        ws[f"A{r}"].font = F_BOLD
        ws[f"B{r}"] = text
        ws[f"B{r}"].font, ws[f"B{r}"].alignment = F_BODY, WRAP
        ws.row_dimensions[r].height = 15 * (len(text) // 105 + 1) + 2
        r += 1
    ws.sheet_view.showGridLines = False


# ---- expected values (computed independently, in SQL) -----------------------

def expected_values(con, last_month_day="2025-12-31"):
    slice_c = f"CAST(signup_date AS DATE) BETWEEN DATE '{SIGNUP_FROM}' AND DATE '{SIGNUP_TO}'"
    q = lambda s: con.execute(s).fetchone()[0]  # noqa: E731
    base = f"""
        WITH c AS (SELECT * FROM stg_customers WHERE {slice_c}),
             o AS (SELECT f.* FROM fct_orders f JOIN c USING (customer_id))
    """
    ana = base + ", a AS (SELECT * FROM o WHERE revenue_recognised AND NOT is_test_account)"
    return {
        "customers": q(base + "SELECT count(*) FROM c"),
        "orders": q(base + "SELECT count(*) FROM o"),
        "test_accounts": q(base + "SELECT count(*) FROM c WHERE is_test_account"),
        "analysis_orders": q(ana + "SELECT count(*) FROM a"),
        "net_revenue": round(float(q(ana + "SELECT sum(net_revenue) FROM a")), 2),
        "statuses": q(base + "SELECT count(DISTINCT status) FROM o"),
        "buyers": q(ana + "SELECT count(DISTINCT customer_id) FROM a"),
        "repeat_buyers": q(ana + """SELECT count(*) FROM (SELECT customer_id FROM a
                                    GROUP BY 1 HAVING count(*) > 1)"""),
        "new_revenue": round(float(q(ana + """
            , f AS (SELECT customer_id, min(order_month) AS fm FROM a GROUP BY 1)
            SELECT sum(a.net_revenue) FROM a JOIN f USING (customer_id)
            WHERE a.order_month = f.fm""")), 2),
        "jan_m1": round(float(q(ana + """
            SELECT count(DISTINCT a.customer_id) * 1.0 /
                   (SELECT count(*) FROM c WHERE signup_month = DATE '2025-01-01'
                    AND NOT is_test_account)
            FROM a WHERE signup_month = DATE '2025-01-01'
              AND date_diff('month', signup_month, order_month) = 1""")), 6),
        "unknown_country": q(base + "SELECT count(*) FROM c WHERE country IS NULL"),
    }


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--db", default=ROOT / "warehouse.duckdb", type=Path)
    ap.add_argument("--out", default=OUT, type=Path)
    args = ap.parse_args()

    con = duckdb.connect(str(args.db))
    customers, orders = load(con)
    expected = expected_values(con)

    wb = Workbook()
    sheet_readme(wb, len(customers), len(orders))
    last_cust = sheet_raw_customers(wb, customers)
    last_ord = sheet_raw_orders(wb, orders)
    sheet_customers(wb, last_cust, last_ord)
    sheet_clean_orders(wb, last_ord, last_cust)
    sheet_monthly(wb, last_ord)
    sheet_cohorts(wb, last_ord, last_cust)
    sheet_segments(wb, last_cust)
    sheet_channels(wb, last_cust)
    sheet_dashboard(wb, last_ord, last_cust)
    sheet_checks(wb, expected, last_ord, last_cust)
    sheet_pivot_guide(wb)

    # Sheet order a reader meets: README, Dashboard, analysis, cleaning, raw.
    order = ["README", "Dashboard", "Monthly", "Cohorts", "Segments", "Channels", "Checks",
             "Customers", "Clean_Orders", "Raw_Customers", "Raw_Orders", "Pivot_Guide"]
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    for name in ("Raw_Customers", "Raw_Orders"):
        wb[name].sheet_properties.tabColor = "7F7F7F"
    wb["Dashboard"].sheet_properties.tabColor = "1F3A5F"
    wb["Checks"].sheet_properties.tabColor = "008000"

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"Wrote {args.out}: {len(customers):,} customers, {len(orders):,} orders")
    print("Expected values:", expected)


if __name__ == "__main__":
    main()
